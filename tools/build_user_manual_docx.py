from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from build_delivery_manual_docx import (
    HEADER_FILL,
    MUTED,
    PRIMARY,
    PRIMARY_DARK,
    add_markdown,
    configure_document,
    set_cell_border,
    set_cell_margins,
    set_run_font,
    shade_cell,
)


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "Manual_Uso_Nomina_Docente.md"
OUTPUT = ROOT / "docs" / "Manual_Uso_Nomina_Docente.docx"


def add_cover(document: Document) -> None:
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(82)
    run = p.add_run("NÓMINA DOCENTE")
    set_run_font(run, 26, True, PRIMARY_DARK)

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Manual de uso del sistema")
    set_run_font(run, 15, False, PRIMARY)

    table = document.add_table(rows=6, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    values = [
        ("Versión", "1.0 operativa"),
        ("Fecha", "9 de mayo de 2026"),
        ("Ambiente", "Producción"),
        ("Web App", "https://nomina-docente-prod.web.app"),
        ("Dominio permitido", "@tecplayacar.edu.mx"),
        ("Uso", "Guía para Admin, Coordinación, Dirección, RH y Finanzas"),
    ]
    for row, (left, right) in zip(table.rows, values):
        for index, text in enumerate((left, right)):
            cell = row.cells[index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, 120, 140, 120, 140)
            set_cell_border(cell)
            if index == 0:
                shade_cell(cell, HEADER_FILL)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(text)
            set_run_font(run, 9.5, index == 0, PRIMARY_DARK if index == 0 else "334155")

    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(20)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Documento operativo para uso diario, revisión de permisos y flujo quincenal.")
    set_run_font(run, 10, False, MUTED)

    document.add_page_break()


def add_footer(document: Document) -> None:
    for section in document.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run("Nómina Docente - Manual de uso")
        set_run_font(run, 8, False, MUTED)


def main() -> None:
    document = Document()
    configure_document(document)
    add_cover(document)
    markdown = SOURCE.read_text(encoding="utf-8")
    add_markdown(document, markdown, SOURCE.parent)
    add_footer(document)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
