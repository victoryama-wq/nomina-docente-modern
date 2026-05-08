from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "Manual_Entrega_Nomina_Docente.md"
OUTPUT = ROOT / "docs" / "Manual_Entrega_Nomina_Docente.docx"

PRIMARY = "0F766E"
PRIMARY_DARK = "0F172A"
MUTED = "64748B"
LIGHT = "E8F3F1"
HEADER_FILL = "DDEFEA"
BORDER = "B7D4CC"


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = BORDER) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=100, bottom=100, end=100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for margin_name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_run_font(run, size: float | None = None, bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = "Aptos"
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    for style_name, size, color in (
        ("Heading 1", 17, PRIMARY_DARK),
        ("Heading 2", 13, PRIMARY),
        ("Heading 3", 11, PRIMARY_DARK),
    ):
        style = document.styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True


def add_cover(document: Document) -> None:
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(80)
    run = p.add_run("NÓMINA DOCENTE")
    set_run_font(run, 26, True, PRIMARY_DARK)

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Manual de entrega técnico-operativo")
    set_run_font(run, 15, False, PRIMARY)

    table = document.add_table(rows=6, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    values = [
        ("Versión", "1.0 productiva"),
        ("Fecha", "8 de mayo de 2026"),
        ("Ambiente", "Producción Google Cloud / Firebase"),
        ("Web App", "https://nomina-docente-prod.web.app"),
        ("Dominio permitido", "@tecplayacar.edu.mx"),
        ("Administrador protegido", "victor.yama@tecplayacar.edu.mx"),
    ]
    for row, (left, right) in zip(table.rows, values):
        for index, text in enumerate((left, right)):
            cell = row.cells[index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, 120, 140, 120, 140)
            set_cell_border(cell)
            if index == 0:
                shade_cell(cell, HEADER_FILL)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(text)
            set_run_font(run, 9.5, index == 0, PRIMARY_DARK if index == 0 else "334155")

    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(20)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Documento de referencia para entrega, operación y mantenimiento del sistema.")
    set_run_font(run, 10, False, MUTED)

    document.add_page_break()


def add_footer(document: Document) -> None:
    for section in document.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run("Nómina Docente - Manual de entrega")
        set_run_font(run, 8, False, MUTED)


def parse_markdown_table(lines: list[str], index: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    while index < len(lines) and lines[index].strip().startswith("|"):
        line = lines[index].strip()
        if re.fullmatch(r"\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?", line):
            index += 1
            continue
        cells = [cell.strip() for cell in line.strip("|").split("|")]
        rows.append(cells)
        index += 1
    return rows, index


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    width = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=width)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for r_idx, row in enumerate(rows):
        for c_idx in range(width):
            text = row[c_idx] if c_idx < len(row) else ""
            cell = table.cell(r_idx, c_idx)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, 110, 120, 110, 120)
            set_cell_border(cell)
            if r_idx == 0:
                shade_cell(cell, HEADER_FILL)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if len(text) > 18 else WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(text.replace("`", ""))
            set_run_font(run, 8.4, r_idx == 0, PRIMARY_DARK if r_idx == 0 else "334155")
    document.add_paragraph()


def add_code_block(document: Document, block: list[str], language: str) -> None:
    title = "Diagrama" if language == "mermaid" else "Bloque técnico"
    p = document.add_paragraph()
    p.paragraph_format.keep_with_next = True
    run = p.add_run(title)
    set_run_font(run, 8.5, True, PRIMARY)
    for line in block:
        p = document.add_paragraph()
        p.style = document.styles["Normal"]
        p.paragraph_format.left_indent = Inches(0.22)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(7.6)
        run.font.color.rgb = RGBColor.from_string("334155")


def add_markdown(document: Document, markdown: str) -> None:
    lines = markdown.splitlines()
    i = 0
    in_code = False
    code_lang = ""
    code_lines: list[str] = []
    skip_until_content = True

    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()

        if skip_until_content:
            if line.startswith("## 1. "):
                skip_until_content = False
            else:
                i += 1
                continue

        if line.startswith("```"):
            if in_code:
                add_code_block(document, code_lines, code_lang)
                code_lines = []
                code_lang = ""
                in_code = False
            else:
                in_code = True
                code_lang = line.strip("`").strip()
            i += 1
            continue

        if in_code:
            code_lines.append(line)
            i += 1
            continue

        if not line.strip() or line.strip() == "---":
            i += 1
            continue

        if line.strip().startswith("|"):
            rows, i = parse_markdown_table(lines, i)
            add_table(document, rows)
            continue

        heading = re.match(r"^(#{1,6})\s+(.*)$", line)
        if heading:
            level = len(heading.group(1))
            text = heading.group(2).strip()
            if level == 2:
                document.add_heading(text, level=1)
            elif level == 3:
                document.add_heading(text, level=2)
            else:
                document.add_heading(text, level=3)
            i += 1
            continue

        bullet = re.match(r"^[-*]\s+(.*)$", line)
        if bullet:
            p = document.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(bullet.group(1).replace("`", ""))
            set_run_font(run, 9.2, False, "334155")
            i += 1
            continue

        numbered = re.match(r"^\d+\.\s+(.*)$", line)
        if numbered:
            p = document.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(numbered.group(1).replace("`", ""))
            set_run_font(run, 9.2, False, "334155")
            i += 1
            continue

        p = document.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        text = line.replace("**", "").replace("`", "")
        run = p.add_run(text)
        set_run_font(run, 9.4, False, "334155")
        i += 1


def main() -> None:
    document = Document()
    configure_document(document)
    add_cover(document)
    markdown = SOURCE.read_text(encoding="utf-8")
    add_markdown(document, markdown)
    add_footer(document)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
